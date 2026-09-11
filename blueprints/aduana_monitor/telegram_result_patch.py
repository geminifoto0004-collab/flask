# -*- coding: utf-8 -*-
"""Readable, paginated Telegram results with on-demand Excel export."""
from __future__ import annotations

import io
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import bot, settings, storage, telegram

_PAGE_SIZE = 4
_BASE_HANDLE_CALLBACK = bot._handle_callback

_EXCEL_COLUMNS = [
    ("n_denuncia", "N° Denuncia"),
    ("emision", "Emisión"),
    ("notificacion", "Notificación"),
    ("doc_aduanero", "Doc. Aduanero"),
    ("art_infraccion", "Art. Infracción"),
    ("infractor", "Infractor"),
    ("multa_max_legal", "Multa Max. Legal"),
    ("multa_c_allan", "Multa c/Allan."),
    ("multa_s_allan", "Multa s/Allan."),
    ("venc_allan", "Venc. Allan."),
    ("venc_recl_junta", "Venc. Recl. Junta"),
    ("audiencia", "Audiencia"),
    ("n_despacho", "N° Despacho"),
    ("aduana_nombre", "Aduana"),
]


def _value(row, key):
    return str((row or {}).get(key) or "").strip()


def _money(value):
    text = str(value or "").strip()
    compact = re.sub(r"\s+", "", text)
    if compact.isdigit():
        try:
            return f"{int(compact):,}".replace(",", ".")
        except Exception:
            pass
    return text


def _result_state(user):
    data = storage.pending_data(user) if user else {}
    rows = data.get("_result_rows")
    if not isinstance(rows, list):
        rows = []
    return {
        "rows": rows,
        "rut": str(data.get("_result_rut") or ""),
        "aduana": str(data.get("_result_aduana") or "7"),
        "all_ok": bool(data.get("_result_all_ok", True)),
    }


def _result_buttons(page, pages):
    rows = []
    if pages > 1:
        nav = []
        if page > 0:
            nav.append(("⬅️ Anterior", f"res_page:{page - 1}"))
        nav.append((f"{page + 1}/{pages}", "res_noop"))
        if page < pages - 1:
            nav.append(("Siguiente ➡️", f"res_page:{page + 1}"))
        rows.append(nav)
    rows.append([("📊 Excel", "res_xlsx"), ("🔄 Otra consulta", "ui_query")])
    return rows


def _render_result_page(user, page=0, message_id=None):
    state = _result_state(user)
    rows = state["rows"]
    rut = state["rut"]
    aduana = state["aduana"]
    all_ok = state["all_ok"]
    if not rows:
        return bot._show_panel(
            user,
            "No hay resultados guardados para mostrar.",
            rows=[[("🔄 Otra consulta", "ui_query")]],
            message_id=message_id,
        )

    pages = max(1, int(math.ceil(len(rows) / _PAGE_SIZE)))
    page = max(0, min(int(page), pages - 1))
    start = page * _PAGE_SIZE
    chunk = rows[start:start + _PAGE_SIZE]
    label = settings.ADUANA_LABELS.get(aduana, aduana)

    infractors = []
    for row in rows:
        name = _value(row, "infractor")
        if name and name not in infractors:
            infractors.append(name)

    warning = "⚠️ Algunos períodos tuvieron error.\n\n" if not all_ok else ""
    lines = [
        warning + f"✅ <b>{len(rows)} resultado(s)</b>",
        f"RUT: <code>{telegram.escape_html(rut)}</code> · Aduana: <b>{telegram.escape_html(label)}</b>",
    ]
    if len(infractors) == 1:
        lines.append(f"👤 {telegram.escape_html(infractors[0])}")

    for offset, row in enumerate(chunk, start=start + 1):
        denuncia = _value(row, "n_denuncia") or "—"
        emision = _value(row, "emision") or "—"
        notificacion = _value(row, "notificacion")
        documento = _value(row, "doc_aduanero") or "—"
        multa = _money(_value(row, "multa_max_legal")) or "—"
        lines.extend([
            "",
            f"<b>{offset}. N° {telegram.escape_html(denuncia)}</b> · {telegram.escape_html(emision)}",
            f"📄 {telegram.escape_html(documento)}",
            f"💰 Multa máx.: <b>{telegram.escape_html(multa)}</b>",
        ])
        if notificacion:
            lines.append(f"🔔 Notificación: {telegram.escape_html(notificacion)}")
        if len(infractors) != 1:
            infractor = _value(row, "infractor")
            if infractor:
                lines.append(f"👤 {telegram.escape_html(infractor)}")

    lines.extend(["", "¿Quieres el archivo completo? Pulsa <b>📊 Excel</b>."])
    return bot._show_panel(
        user,
        "\n".join(lines),
        rows=_result_buttons(page, pages),
        message_id=message_id,
    )


def _build_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Aduana"
    ws.append([label for _key, label in _EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([str((row or {}).get(key) or "") for key, _label in _EXCEL_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, (key, label) in enumerate(_EXCEL_COLUMNS, start=1):
        sample = [label] + [str((row or {}).get(key) or "") for row in rows[:200]]
        width = min(42, max(10, max(len(value) for value in sample) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _send_excel(user, message_id=None):
    state = _result_state(user)
    rows = state["rows"]
    if not rows:
        return _render_result_page(user, 0, message_id=message_id)
    rut = state["rut"]
    aduana = state["aduana"]
    label = settings.ADUANA_LABELS.get(aduana, aduana)
    safe_label = re.sub(r"[^A-Za-z0-9_-]+", "_", str(label or "ADUANA"))
    safe_rut = re.sub(r"[^0-9Kk-]+", "", rut)
    telegram.send_document(
        user["chat_id"],
        f"ADUANA_{safe_label}_{safe_rut}.xlsx",
        _build_xlsx(rows),
        f"{len(rows)} resultado(s) · {rut} · {label}",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    # Keep the result panel in place; the Excel file appears as a separate file
    # message because Telegram documents cannot replace an inline result panel.
    return None


def _send_results(user_id, chat_id, message_id, rows, rut, aduana, all_ok):
    user = storage.get_user(user_id)
    if not user:
        return

    if not rows and not all_ok:
        bot._set_state(user_id, "IDLE", {}, panel_id=message_id or None)
        return bot._show_panel(
            storage.get_user(user_id) or user,
            "⚠️ <b>Consulta incompleta</b>\n\n"
            "Aduana no respondió correctamente en uno o más períodos. "
            "No podemos afirmar que no existan resultados.",
            rows=[[("🔄 Intentar otra consulta", "ui_query")]],
            message_id=message_id or None,
        )

    if not rows:
        bot._set_state(user_id, "IDLE", {}, panel_id=message_id or None)
        return bot._show_panel(
            storage.get_user(user_id) or user,
            "Sin resultados para ese período.",
            rows=[[("🔄 Otra consulta", "ui_query")]],
            message_id=message_id or None,
        )

    # Persist the current result set in the user's UI state so every record can
    # be browsed in Telegram page by page and Excel can be requested later.
    bot._set_state(
        user_id,
        "RESULT",
        {
            "_result_rows": rows,
            "_result_rut": rut,
            "_result_aduana": aduana,
            "_result_all_ok": bool(all_ok),
        },
        panel_id=message_id or None,
    )
    return _render_result_page(storage.get_user(user_id) or user, 0, message_id=message_id or None)


def _handle_callback(user, data, message_id=None):
    if data == "res_noop":
        return None
    if data.startswith("res_page:"):
        try:
            page = int(data.split(":", 1)[1])
        except Exception:
            page = 0
        return _render_result_page(user, page, message_id=message_id)
    if data == "res_xlsx":
        return _send_excel(user, message_id=message_id)
    return _BASE_HANDLE_CALLBACK(user, data, message_id)


def install():
    if getattr(bot, "_telegram_result_patch_installed", False):
        return
    bot._send_results = _send_results
    bot._handle_callback = _handle_callback
    bot._telegram_result_patch_installed = True


install()
