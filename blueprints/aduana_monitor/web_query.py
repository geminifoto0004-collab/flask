# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import json
import re
import time
from datetime import date, datetime

import requests
from flask import (
    abort,
    jsonify,
    redirect,
    render_template,
    render_template_string,
    request,
    send_file,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import aduana_bp, collector_queue, query, settings


def _normalize_rut(value):
    compact = re.sub(r"[^0-9Kk]", "", str(value or "")).upper()
    if len(compact) not in (8, 9) or not compact[:-1].isdigit():
        raise ValueError("RUT inválido")
    return compact[:-1] + "-" + compact[-1]


WEB_COLUMNS = [
    ("n_denuncia", "N° Denuncia"),
    ("emision", "Emisión"),
    ("notificacion", "Notificación"),
    ("doc_aduanero", "Doc. Aduanero"),
    ("art_infraccion", "Art. Infracción"),
    ("infractor", "Infractor"),
    ("multa_max_legal", "Multa Max. Legal"),
    ("multa_c_allan", "Multa c/Allan."),
    ("multa_s_allan", "Multa s/Allan."),
    ("venc_allan", "Venc. Allan."),
    ("venc_recl_junta", "Venc. Recl. Junta"),
    ("audiencia", "Audiencia"),
    ("n_despacho", "N° Despacho"),
    ("aduana_nombre", "Aduana"),
]


def _failure_examples(logs, limit=4):
    examples = []
    seen = set()
    for item in logs or []:
        if str(item.get("estado") or "").upper() == "OK":
            continue
        error = " ".join(str(item.get("error") or "sin detalle").split())
        key = (str(item.get("estado") or ""), str(item.get("phase") or ""), error)
        if key in seen:
            continue
        seen.add(key)
        examples.append({
            "estado": item.get("estado") or "ERROR",
            "error": error,
            "desde": item.get("desde") or "",
            "hasta": item.get("hasta") or "",
            "worker": item.get("worker") or "",
            "phase": item.get("phase") or "",
        })
        if len(examples) >= limit:
            break
    return examples


def _year_choices():
    current_year = date.today().year
    primary_years = [current_year, current_year - 1, current_year - 2]
    older_years = list(range(current_year - 3, settings.UNLIMITED_START_YEAR - 1, -1))
    return current_year, primary_years, older_years, primary_years + older_years


def _periods_for_years(years):
    today = date.today()
    periods = []
    for year in sorted({int(y) for y in years}, reverse=True):
        start = date(year, 1, 1)
        end = today if year == today.year else date(year, 12, 31)
        periods.extend(query.month_chunks_for_range(start, end))
    return periods


def _render_query(
    *,
    selected_years,
    aduana="7",
    rut="",
    rows=None,
    logs=None,
    searched=False,
    all_ok=True,
    message="",
    elapsed=None,
    export_url="",
):
    _current_year, primary_years, older_years, _available = _year_choices()
    rows = rows or []
    logs = logs or []
    return render_template(
        "aduana_admin/query.html",
        primary_years=primary_years,
        older_years=older_years,
        selected_years=selected_years,
        aduanas=settings.ADUANAS,
        aduana=aduana,
        aduana_label=settings.ADUANA_LABELS.get(aduana, aduana),
        rut=rut,
        rows=rows,
        logs=logs,
        failure_examples=_failure_examples(logs),
        columns=WEB_COLUMNS,
        searched=searched,
        all_ok=all_ok,
        message=message,
        elapsed=elapsed,
        export_url=export_url,
    )


def _public_job(job_id, request_id):
    job = collector_queue.get_job(job_id)
    if not job or str(job.get("request_id") or "") != str(request_id or ""):
        abort(404)
    return job


def _job_selected_years(job):
    try:
        payload = json.loads(job.get("payload_json") or "{}")
    except Exception:
        payload = {}
    years = set()
    for item in payload.get("periods") or []:
        raw = str(item.get("start") or "")
        try:
            years.add(int(raw[:4]))
        except Exception:
            continue
    return sorted(years, reverse=True)


def _parse_db_datetime(value):
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, str) and value:
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
        except Exception:
            return None
    return None


def _job_elapsed(job):
    started = _parse_db_datetime(job.get("created_at"))
    ended = _parse_db_datetime(job.get("completed_at")) or datetime.now()
    if not started:
        return None
    try:
        return max(0.0, round((ended - started).total_seconds(), 2))
    except Exception:
        return None


def _job_result_payload(job, request_id):
    status = str(job.get("status") or "UNKNOWN").upper()
    selected_years = _job_selected_years(job)
    aduana = str(job.get("aduana") or "7")
    rut = str(job.get("rut") or "")
    elapsed = _job_elapsed(job)

    if status in ("PENDING", "RUNNING"):
        return {
            "ready": False,
            "status": status,
            "selected_years": selected_years,
            "aduana": aduana,
            "aduana_label": settings.ADUANA_LABELS.get(aduana, aduana),
            "rut": rut,
            "elapsed": elapsed,
        }

    if status == "DONE":
        try:
            result = json.loads(job.get("result_json") or "{}")
        except Exception as exc:
            result = {}
            status = "FAILED"
            job["error_message"] = f"Invalid worker result: {exc}"

    if status == "DONE":
        rows = result.get("rows") if isinstance(result.get("rows"), list) else []
        logs = result.get("logs") if isinstance(result.get("logs"), list) else []
        all_ok = bool(result.get("all_ok"))
        failed = sum(1 for item in logs if item.get("estado") != "OK")
        message = (
            f"查詢完成，共 {len(rows)} 筆。"
            if all_ok
            else f"查詢完成，共 {len(rows)} 筆；{failed} 個月份查詢失敗。"
        )
    else:
        rows = []
        error = str(job.get("error_message") or f"Worker job {status}")
        logs = [{
            "worker": job.get("worker_id") or "REMOTE",
            "desde": "",
            "hasta": "",
            "estado": "REQUEST_FAILED",
            "resultado": "",
            "filas": 0,
            "segundos": 0,
            "phase": "WORKER",
            "error": error,
        }]
        all_ok = False
        message = f"查詢失敗：{error}"

    export_url = ""
    if rows:
        export_url = url_for(
            ".public_query_export",
            job_id=int(job["id"]),
            request_id=request_id,
        )

    return {
        "ready": True,
        "status": status,
        "selected_years": selected_years,
        "aduana": aduana,
        "aduana_label": settings.ADUANA_LABELS.get(aduana, aduana),
        "rut": rut,
        "rows": rows,
        "logs": logs,
        "failure_examples": _failure_examples(logs),
        "all_ok": all_ok,
        "message": message,
        "elapsed": elapsed,
        "export_url": export_url,
    }


_WAIT_HTML = """
<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Consulta de Denuncias</title>
<style>
body{margin:0;background:#f4f6f8;color:#18212f;font-family:Inter,'Segoe UI',Arial,'Microsoft JhengHei',sans-serif}
.wrap{max-width:820px;margin:70px auto;padding:0 18px}.card{background:#fff;border:1px solid #dfe5eb;border-radius:14px;padding:26px;box-shadow:0 3px 14px rgba(28,44,64,.06)}
.spinner{width:34px;height:34px;border:4px solid #dce5ec;border-top-color:#174a73;border-radius:50%;animation:r 1s linear infinite;margin-bottom:18px}@keyframes r{to{transform:rotate(360deg)}}
h2{margin:0 0 10px;font-size:20px}.muted{color:#6f7c89;font-size:13px}.state{margin-top:18px;padding:10px 12px;background:#f5f8fa;border-radius:8px;font-size:13px}
</style>
</head>
<body><div class="wrap"><div class="card">
<div class="spinner"></div><h2>正在查詢 Aduana…</h2>
<div class="muted">Windows Worker 正在處理。</div>
<div class="state" id="state">狀態：等待 Windows Worker…</div>
</div></div>
<script>
const statusUrl={{ status_url|tojson }};
const resultUrl={{ result_url|tojson }};
async function poll(){
  try{
    const r=await fetch(statusUrl,{cache:'no-store'});
    const d=await r.json();
    const el=document.getElementById('state');
    if(d.status==='PENDING') el.textContent='狀態：等待 Windows Worker…';
    else if(d.status==='RUNNING') el.textContent='狀態：Windows Worker 已接到工作，正在查詢…';
    else if(['DONE','FAILED','EXPIRED'].includes(d.status)){window.location.replace(resultUrl);return;}
  }catch(e){document.getElementById('state').textContent='Render 暫時沒有回應，正在重試…';}
  setTimeout(poll,900);
}
poll();
</script></body></html>
"""


@aduana_bp.route("/aduana", methods=["GET", "POST"])
@aduana_bp.route(f"{settings.ADMIN_PREFIX}/query", methods=["GET", "POST"])
def admin_query():
    current_year, _primary_years, _older_years, available_years = _year_choices()

    if request.method == "GET":
        return _render_query(selected_years=[current_year])

    ajax = request.headers.get("X-Requested-With") == "AduanaAjax"
    selected_years = []
    aduana = str(request.form.get("aduana") or "7").strip()
    rut_input = str(request.form.get("rut") or "").strip()
    for raw in request.form.getlist("years"):
        try:
            year = int(raw)
        except Exception:
            continue
        if year in available_years:
            selected_years.append(year)
    selected_years = sorted(set(selected_years), reverse=True)

    try:
        if not selected_years:
            raise ValueError("至少選擇一個年份")
        if aduana not in settings.ADUANA_CODES:
            raise ValueError("請選擇 Aduana")
        rut = _normalize_rut(rut_input)

        if settings.WORKER_ENABLED:
            periods = _periods_for_years(selected_years)
            job_id = collector_queue.create_job(
                periods,
                aduana,
                rut,
                max_workers=settings.PERIOD_WORKERS,
            )
            job = collector_queue.get_job(job_id)
            if not job:
                raise RuntimeError("Worker job could not be created")
            request_id = str(job.get("request_id") or "")
            if ajax:
                return jsonify({
                    "ok": True,
                    "job_id": int(job_id),
                    "request_id": request_id,
                    "status_url": url_for(
                        ".public_query_status",
                        job_id=job_id,
                        request_id=request_id,
                    ),
                    "result_api_url": url_for(
                        ".public_query_result_api",
                        job_id=job_id,
                        request_id=request_id,
                    ),
                })
            return redirect(
                url_for(
                    ".public_query_wait",
                    job_id=job_id,
                    request_id=request_id,
                ),
                code=303,
            )

        started = time.perf_counter()
        rows, logs, all_ok = query.query_years(selected_years, aduana, rut)
        elapsed = round(time.perf_counter() - started, 2)
        failed = sum(1 for item in logs if item.get("estado") != "OK")
        message = (
            f"查詢完成，共 {len(rows)} 筆。"
            if all_ok
            else f"查詢完成，共 {len(rows)} 筆；{failed} 個月份查詢失敗。"
        )
        if ajax:
            return jsonify({
                "ok": True,
                "direct_result": {
                    "ready": True,
                    "status": "DONE" if all_ok else "FAILED",
                    "selected_years": selected_years,
                    "aduana": aduana,
                    "aduana_label": settings.ADUANA_LABELS.get(aduana, aduana),
                    "rut": rut,
                    "rows": rows,
                    "logs": logs,
                    "failure_examples": _failure_examples(logs),
                    "all_ok": all_ok,
                    "message": message,
                    "elapsed": elapsed,
                    "export_url": "",
                },
            })
        return _render_query(
            selected_years=selected_years,
            aduana=aduana,
            rut=rut,
            rows=rows,
            logs=logs,
            searched=True,
            all_ok=all_ok,
            message=message,
            elapsed=elapsed,
        )
    except Exception as exc:
        if ajax:
            return jsonify({"ok": False, "error": str(exc)}), 400
        return _render_query(
            selected_years=selected_years or [current_year],
            aduana=aduana,
            rut=rut_input,
            searched=True,
            all_ok=False,
            message=str(exc),
        )


@aduana_bp.route("/aduana/job/<int:job_id>/<request_id>", methods=["GET"])
def public_query_wait(job_id, request_id):
    job = _public_job(job_id, request_id)
    status = str(job.get("status") or "").upper()
    if status in ("DONE", "FAILED", "EXPIRED"):
        return redirect(
            url_for(".public_query_result", job_id=job_id, request_id=request_id),
            code=303,
        )
    return render_template_string(
        _WAIT_HTML,
        status_url=url_for(
            ".public_query_status",
            job_id=job_id,
            request_id=request_id,
        ),
        result_url=url_for(
            ".public_query_result",
            job_id=job_id,
            request_id=request_id,
        ),
    )


@aduana_bp.route("/api/aduana/public-query/<int:job_id>/<request_id>", methods=["GET"])
def public_query_status(job_id, request_id):
    job = _public_job(job_id, request_id)
    return jsonify({
        "ok": True,
        "status": str(job.get("status") or "UNKNOWN").upper(),
        "worker_id": job.get("worker_id"),
        "attempts": int(job.get("attempts") or 0),
    })


@aduana_bp.route("/api/aduana/public-query/<int:job_id>/<request_id>/result", methods=["GET"])
def public_query_result_api(job_id, request_id):
    job = _public_job(job_id, request_id)
    payload = _job_result_payload(job, request_id)
    return jsonify({"ok": True, **payload}), (200 if payload.get("ready") else 202)


@aduana_bp.route("/aduana/result/<int:job_id>/<request_id>", methods=["GET"])
def public_query_result(job_id, request_id):
    job = _public_job(job_id, request_id)
    payload = _job_result_payload(job, request_id)
    if not payload.get("ready"):
        return redirect(
            url_for(".public_query_wait", job_id=job_id, request_id=request_id),
            code=303,
        )
    return _render_query(
        selected_years=payload["selected_years"],
        aduana=payload["aduana"],
        rut=payload["rut"],
        rows=payload["rows"],
        logs=payload["logs"],
        searched=True,
        all_ok=payload["all_ok"],
        message=payload["message"],
        elapsed=payload["elapsed"],
        export_url=payload.get("export_url") or "",
    )


@aduana_bp.route("/aduana/export/<int:job_id>/<request_id>.xlsx", methods=["GET"])
def public_query_export(job_id, request_id):
    job = _public_job(job_id, request_id)
    payload = _job_result_payload(job, request_id)
    if not payload.get("ready"):
        return jsonify({"ok": False, "error": "query still running"}), 409
    rows = payload.get("rows") or []
    if not rows:
        return jsonify({"ok": False, "error": "no rows to export"}), 404

    wb = Workbook()
    ws = wb.active
    ws.title = "Aduana"
    headers = [label for _key, label in WEB_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([str(row.get(key, "") or "") for key, _label in WEB_COLUMNS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, (key, label) in enumerate(WEB_COLUMNS, start=1):
        sample = [label] + [str(row.get(key, "") or "") for row in rows[:200]]
        width = min(42, max(10, max(len(value) for value in sample) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    years = "-".join(str(y) for y in payload.get("selected_years") or []) or "consulta"
    label = re.sub(r"[^A-Za-z0-9_-]+", "_", str(payload.get("aduana_label") or "ADUANA"))
    rut = re.sub(r"[^0-9Kk-]+", "", str(payload.get("rut") or ""))
    filename = f"ADUANA_{label}_{rut}_{years}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# Temporary, no-signup proxy connectivity probe.
# IMPORTANT: this only requests the public Aduana landing page; it never sends a RUT.
@aduana_bp.route("/aduana/proxy-test-public", methods=["GET"])
def proxy_test_public():
    proxy_url = "http://45.225.204.11:999"
    started = time.perf_counter()
    try:
        response = requests.get(
            settings.BASE_URL,
            headers=settings.HEADERS,
            proxies={"http": proxy_url, "https": proxy_url},
            timeout=6,
            allow_redirects=True,
        )
        elapsed = round(time.perf_counter() - started, 2)
        text = response.text or ""
        return jsonify({
            "ok": response.status_code == 200 and "P1_FECHA_DESDE" in text,
            "status_code": response.status_code,
            "elapsed_seconds": elapsed,
            "final_url": response.url,
            "aduana_form_found": "P1_FECHA_DESDE" in text,
            "response_bytes": len(response.content or b""),
            "note": "Public proxy probe only; no RUT was sent.",
        })
    except Exception as exc:
        return jsonify({
            "ok": False,
            "elapsed_seconds": round(time.perf_counter() - started, 2),
            "error": f"{type(exc).__name__}: {exc}",
            "note": "Public proxy probe only; no RUT was sent.",
        }), 502
